# -*- coding: utf-8 -*-
"""
ملف الاختبار الشامل لجميع وحدات ووظائف النظام
"""

import os
import sys
import shutil

# ضبط ترميز الطرفية على UTF-8 لطباعة الحروف العربية في ويندوز
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

# إضافة مسار employee_system
current_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(current_dir, "employee_system"))

from database import Database
from file_operations import (
    export_employees_to_excel,
    parse_and_validate_excel_file,
    execute_import_commit,
    export_word_employee_cards,
    export_word_employee_table,
    backup_database_to_file,
    restore_database_from_file
)
import openpyxl

TEST_DIR = os.path.join(current_dir, "test_output")
os.makedirs(TEST_DIR, exist_ok=True)
TEST_DB_PATH = os.path.join(TEST_DIR, "test_database.db")

def run_tests():
    print("=== بدء الفحص الآلي الشامل لنظام سجل الموظفين ===")

    # 1. اختبار قاعدة البيانات والتهيئة الأولية
    print("\n1. اختبار تهيئة قاعدة البيانات والتلقيم الافتراضي:")
    if os.path.exists(TEST_DB_PATH):
        os.remove(TEST_DB_PATH)

    db = Database(db_path=TEST_DB_PATH)
    users = db.get_users()
    statuses = db.get_statuses()
    
    assert len(users) == 2, f"Expected 2 users, got {len(users)}"
    assert any(u["role"] == "admin" for u in users), "Admin user not found"
    assert any(u["role"] == "staff" for u in users), "Staff user not found"
    print(f"  [+] تم تلقيم المستخدمين بنجاح: {[u['display_name'] for u in users]}")

    assert len(statuses) >= 5, f"Expected at least 5 statuses, got {len(statuses)}"
    status_names = [s["name"] for s in statuses]
    assert "على رأس العمل" in status_names
    assert "متقاعد" in status_names
    print(f"  [+] تم تلقيم الحالات الوظيفية بنجاح: {status_names}")

    # 2. اختبار المصادقة
    print("\n2. اختبار مصادقة المستخدمين:")
    admin_auth = db.authenticate_user(users[0]["id"], "1234")
    assert admin_auth is not None and admin_auth["role"] == "admin", "Admin auth failed"
    wrong_auth = db.authenticate_user(users[0]["id"], "9999")
    assert wrong_auth is None, "Wrong PIN should fail"
    print("  [+] نجحت عمليات المصادقة والتحقق من رمز PIN")

    # 3. اختبار الحقول الديناميكية
    print("\n3. اختبار إدارة الحقول الديناميكية:")
    f1 = db.add_field("phone", "رقم الهاتف المحمول", display_order=1)
    f2 = db.add_field("address", "عنوان السكن", display_order=2)
    active_fields = db.get_fields(active_only=True)
    assert len(active_fields) == 2
    print(f"  [+] تمت إضافة الحقول الديناميكية بنجاح: {[f['label'] for f in active_fields]}")

    # 4. اختبار إضافة الموظفين والتحقق من الرقم الوطني
    print("\n4. اختبار إضافة موظف والتحقق من منع تكرار الرقم الوطني:")
    emp1_data = {
        "national_id": "119850123456",
        "full_name": "أحمد محمود حسن",
        "status_id": statuses[0]["id"],
        "hire_date": "2015-03-01",
        "department": "الشؤون الإدارية",
        "current_grade": "الأولى",
        "grade_date": "2021-01-01",
        "education_level": "بكالوريوس",
        "specialization": "إدارة أعمال"
    }
    emp1_custom = {"phone": "0912345678", "address": "طرابلس"}
    res1 = db.save_employee(emp1_data, emp1_custom)
    emp1_id = res1["id"]
    assert emp1_id > 0

    # فحص الرقم الوطني
    exists, existing_info = db.check_national_id_exists("119850123456")
    assert exists is True
    assert existing_info["full_name"] == "أحمد محمود حسن"
    print(f"  [+] كشف التكرار للرقم الوطني يعمل بنجاح: {existing_info['full_name']}")

    # محاولة إضافة موظف بنفس الرقم الوطني يجب أن تفشل
    duplicate_failed = False
    try:
        emp2_data = dict(emp1_data)
        emp2_data["full_name"] = "شخص آخر"
        db.save_employee(emp2_data)
    except ValueError as ve:
        duplicate_failed = True
        print(f"  [+] تم منع التكرار بنجاح: {str(ve)}")
    assert duplicate_failed, "Duplicate national_id should raise ValueError"

    # إضافة موظف ثانٍ
    emp2_data = {
        "national_id": "219900987654",
        "full_name": "فاطمة علي محمد",
        "status_id": statuses[0]["id"],
        "hire_date": "2019-07-15",
        "department": "المكتب الفني",
        "current_grade": "الثالثة",
        "grade_date": "2023-01-01",
        "education_level": "ماجستير",
        "specialization": "تقنية معلومات"
    }
    emp2_custom = {"phone": "0923456789", "address": "بنغازي"}
    res2 = db.save_employee(emp2_data, emp2_custom)
    emp2_id = res2["id"]
    print(f"  [+] تمت إضافة الموظفين بنجاح: المعرفات ({emp1_id}, {emp2_id})")

    # جلب تفاصيل موظف مع الحقول الديناميكية
    emp1_fetched = db.get_employee(emp1_id)
    assert emp1_fetched["custom_values"]["phone"] == "0912345678"
    assert emp1_fetched["custom_values"]["address"] == "طرابلس"
    print(f"  [+] تم جلب القيم الديناميكية للموظف بنجاح: {emp1_fetched['custom_values']}")

    # 5. اختبار البحث والفلترة
    print("\n5. اختبار البحث والفلترة:")
    search_res = db.get_employees(search_name="فاطمة")
    assert len(search_res) == 1 and search_res[0]["id"] == emp2_id
    search_nat = db.get_employees(search_national_id="119850123456")
    assert len(search_nat) == 1 and search_nat[0]["id"] == emp1_id
    print("  [+] نجحت عمليات البحث بالاسم وبالرقم الوطني")

    # 6. اختبار تصدير الإكسل
    print("\n6. اختبار تصدير الإكسل (openpyxl):")
    all_emps = db.get_employees()
    excel_export_path = os.path.join(TEST_DIR, "exported_employees.xlsx")
    export_employees_to_excel(all_emps, active_fields, excel_export_path)
    assert os.path.exists(excel_export_path)
    wb_test = openpyxl.load_workbook(excel_export_path)
    ws_test = wb_test.active
    assert ws_test.views.sheetView[0].rightToLeft is True, "Sheet must be RTL"
    assert ws_test.cell(row=1, column=1).value == "الرقم الوطني"
    print(f"  [+] تم إنشاء وتنسيق ملف الإكسل (RTL) بنجاح: {excel_export_path}")

    # 7. اختبار معالج استيراد الإكسل
    print("\n7. اختبار معالج استيراد الإكسل وفض التعارض:")
    
    # 7.أ: اختبار الإيقاف الفوري لملف به خطأ (عدم وجود الرقم الوطني أو الاسم)
    invalid_excel_path = os.path.join(TEST_DIR, "invalid_sample.xlsx")
    wb_err = openpyxl.Workbook()
    ws_err = wb_err.active
    ws_err.append(["الرقم الوطني", "الاسم الكامل", "الحالة الوظيفية"])
    ws_err.append(["111222333444", "موظف صالح", "على رأس العمل"])
    ws_err.append(["", "موظف مفقود الرقم الوطني", "على رأس العمل"]) # سطر خطأ رقم 3
    wb_err.save(invalid_excel_path)

    parse_err_res = parse_and_validate_excel_file(invalid_excel_path, db)
    assert parse_err_res["success"] is False
    assert "السطر رقم (3)" in parse_err_res["error"]
    assert "توقفت العملية" in parse_err_res["error"]
    print(f"  [+] تم التحقق الصارم والإيقاف الفوري بالرسالة المطلوبة: {parse_err_res['error']}")

    # 7.ب: اختبار كشف التعارض وفضه
    import_test_path = os.path.join(TEST_DIR, "import_with_conflict.xlsx")
    wb_conf = openpyxl.Workbook()
    ws_conf = wb_conf.active
    ws_conf.append(["الرقم الوطني", "الاسم الكامل", "الحالة الوظيفية", "القسم"])
    # صف متعارض (نفس الرقم الوطني للموظف 1 ولكن باسم جديد)
    ws_conf.append(["119850123456", "أحمد محمود حسن المعدل", "على رأس العمل", "إدارة المشتريات"])
    # صف جديد
    ws_conf.append(["319950555666", "سارة إبراهيم القاضي", "على رأس العمل", "الشؤون القانونية"])
    wb_conf.save(import_test_path)

    parse_conf_res = parse_and_validate_excel_file(import_test_path, db)
    assert parse_conf_res["success"] is True
    assert len(parse_conf_res["conflicts"]) == 1
    assert len(parse_conf_res["new_records"]) == 1
    conflict_item = parse_conf_res["conflicts"][0]
    assert conflict_item["national_id"] == "119850123456"
    assert conflict_item["existing_name"] == "أحمد محمود حسن"
    assert conflict_item["new_name"] == "أحمد محمود حسن المعدل"
    print(f"  [+] تم اكتشاف التعارض بدقة: مسجل باسم [{conflict_item['existing_name']}] والوارد في الملف [{conflict_item['new_name']}]")

    # تنفيذ الاستيراد مع اختيار تحديث المتعارض
    commit_res = execute_import_commit(
        parse_conf_res["new_records"],
        [{"incoming_record": conflict_item["incoming_record"], "existing_id": conflict_item["existing_id"], "action": "update"}],
        db
    )
    assert commit_res["success"] is True
    assert commit_res["inserted_count"] == 1
    assert commit_res["updated_count"] == 1
    
    # التأكد من التحديث في قاعدة البيانات
    emp1_updated = db.get_employee(emp1_id)
    assert emp1_updated["full_name"] == "أحمد محمود حسن المعدل"
    assert emp1_updated["department"] == "إدارة المشتريات"
    print(f"  [+] تم تنفيذ الاستيراد والتحديث بنجاح: {commit_res['message']}")

    # 8. اختبار تصدير مستندات الوورد (.docx)
    print("\n8. اختبار تصدير مستندات الوورد (python-docx):")
    # 8.أ: بطاقات موظفين مستقلة
    word_cards_path = os.path.join(TEST_DIR, "employee_cards.docx")
    all_emps_current = db.get_employees()
    export_word_employee_cards(all_emps_current, active_fields, None, word_cards_path)
    assert os.path.exists(word_cards_path)
    print(f"  [+] تم توليد مستند بطاقات الموظفين بنجاح: {word_cards_path}")

    # 8.ب: كشف جدولي مجمع
    word_table_path = os.path.join(TEST_DIR, "employee_table.docx")
    export_word_employee_table(all_emps_current, active_fields, None, word_table_path)
    assert os.path.exists(word_table_path)
    print(f"  [+] تم توليد مستند كشف الموظفين المجمع (Landscape) بنجاح: {word_table_path}")

    # 9. اختبار النسخ الاحتياطي والاسترجاع
    print("\n9. اختبار النسخ الاحتياطي والاسترجاع:")
    backup_path = os.path.join(TEST_DIR, "backup_test.db")
    bak_res = backup_database_to_file(TEST_DB_PATH, backup_path)
    assert bak_res["success"] is True
    assert os.path.exists(backup_path)
    print(f"  [+] تم أخذ النسخة الاحتياطية بنجاح: {backup_path}")

    # حذف سجل للتأكد من الاسترجاع
    db.delete_employee_permanent(emp1_id)
    assert db.get_employee(emp1_id) is None
    print("  [+] تم حذف موظف مؤقتاً لاختبار الاسترجاع")

    # استرجاع النسخة
    rest_res = restore_database_from_file(backup_path, TEST_DB_PATH)
    assert rest_res["success"] is True
    db.init_db()
    restored_emp1 = db.get_employee(emp1_id)
    assert restored_emp1 is not None
    assert restored_emp1["full_name"] == "أحمد محمود حسن المعدل"
    print(f"  [+] تم استرجاع قاعدة البيانات وعودة الموظف المحذوف بنجاح: {restored_emp1['full_name']}")

    print("\n========================================================")
    print("  جميع الاختبارات المنهجية (9/9) تمت بنجاح وبأعلى معايير الجودة!")
    print("========================================================")

if __name__ == "__main__":
    run_tests()
