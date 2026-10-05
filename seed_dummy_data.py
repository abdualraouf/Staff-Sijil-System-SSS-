# -*- coding: utf-8 -*-
"""
أداة تلقيم بيانات تجريبية واقعية لنظام سجل الموظفين (Staff Sijil System)
تقوم بإنشاء حقول ديناميكية نموذجية وتلقيم 20 موظفاً ببيانات عربية/ليبية إدارية متكاملة
كما تقوم بإنشاء ملف إكسل نموذجي لاختبار ميزة الاستيراد وفض التعارض
"""

import os
import sys
from typing import Dict, Any, List
from database import Database, get_default_db_path

def seed_database(db: Database) -> None:
    print("=== بدء عملية تلقيم البيانات التجريبية ===")
    
    # 1. إضافة الحقول الديناميكية إذا لم تكن موجودة
    existing_fields = {f["field_key"]: f for f in db.get_fields(active_only=False)}
    sample_fields = [
        ("phone_number", "رقم الهاتف المحمول", 1),
        ("residential_address", "عنوان السكن الحالي", 2),
        ("blood_type", "فصيلة الدم", 3),
    ]
    
    for key, label, order in sample_fields:
        if key not in existing_fields:
            db.add_field(key, label, order)
            print(f"[+] تم إنشاء الحقل الديناميكي: {label} ({key})")

    # جلب الحالات المتاحة
    statuses = {s["name"]: s["id"] for s in db.get_statuses()}
    status_active = statuses.get("على رأس العمل", 1)
    status_leave = statuses.get("إجازة خاصة", 2)
    status_delegated = statuses.get("منتدب", 3)
    status_absent = statuses.get("منقطع", 4)
    status_retired = statuses.get("متقاعد", 5)

    # 2. قائمة الموظفين العشرين
    dummy_employees: List[Dict[str, Any]] = [
        {
            "national_id": "119840012345",
            "full_name": "طارق عبد السلام محمد الزوي",
            "status_id": status_active,
            "hire_date": "2008-04-15",
            "department": "إدارة تقنية المعلومات والتوثيق",
            "current_grade": "الدرجة 11",
            "grade_date": "2021-06-01",
            "education_level": "بكالوريوس",
            "specialization": "هندسة برمجيات",
            "custom_fields": {
                "phone_number": "091-3456789",
                "residential_address": "طرابلس - حي الأندلس",
                "blood_type": "O+"
            }
        },
        {
            "national_id": "219890045678",
            "full_name": "فاطمة مصطفى أحمد الفيتوري",
            "status_id": status_active,
            "hire_date": "2013-09-01",
            "department": "مكتب الموارد البشرية",
            "current_grade": "الدرجة 9",
            "grade_date": "2022-01-15",
            "education_level": "ماجستير",
            "specialization": "إدارة موارد بشرية",
            "custom_fields": {
                "phone_number": "092-5678901",
                "residential_address": "طرابلس - بن عاشور",
                "blood_type": "A+"
            }
        },
        {
            "national_id": "119760078901",
            "full_name": "عمران سعد عبد الله الورفلي",
            "status_id": status_active,
            "hire_date": "2002-11-20",
            "department": "إدارة الشؤون الإدارية والمالية",
            "current_grade": "الدرجة 13",
            "grade_date": "2019-07-01",
            "education_level": "بكالوريوس",
            "specialization": "محاسبة مالية",
            "custom_fields": {
                "phone_number": "091-2345678",
                "residential_address": "طرابلس - زاوية الدهماني",
                "blood_type": "B+"
            }
        },
        {
            "national_id": "219920034567",
            "full_name": "أسماء خليفة سالم المقرحي",
            "status_id": status_active,
            "hire_date": "2016-03-10",
            "department": "مكتب الشؤون القانونية",
            "current_grade": "الدرجة 8",
            "grade_date": "2023-04-01",
            "education_level": "ليسانس",
            "specialization": "قانون عام",
            "custom_fields": {
                "phone_number": "094-7890123",
                "residential_address": "طرابلس - السياحية",
                "blood_type": "AB+"
            }
        },
        {
            "national_id": "119880056789",
            "full_name": "يوسف إبراهيم علي الترهوني",
            "status_id": status_active,
            "hire_date": "2012-07-15",
            "department": "المكتب الفني والمشروعات",
            "current_grade": "الدرجة 10",
            "grade_date": "2021-12-01",
            "education_level": "بكالوريوس",
            "specialization": "هندسة مدنية",
            "custom_fields": {
                "phone_number": "091-8901234",
                "residential_address": "طرابلس - الدريبي",
                "blood_type": "O-"
            }
        },
        {
            "national_id": "119950067890",
            "full_name": "المهدي سالم فرج المجبري",
            "status_id": status_active,
            "hire_date": "2018-10-01",
            "department": "قسم المشتريات والمخازن",
            "current_grade": "الدرجة 7",
            "grade_date": "2022-11-01",
            "education_level": "دبلوم عالي",
            "specialization": "إدارة مخازن وإمداد",
            "custom_fields": {
                "phone_number": "092-1234567",
                "residential_address": "طرابلس - عين زارة",
                "blood_type": "A-"
            }
        },
        {
            "national_id": "219900012344",
            "full_name": "مريم عادل نوري القماطي",
            "status_id": status_active,
            "hire_date": "2014-05-18",
            "department": "قسم العلاقات العامة والإعلام",
            "current_grade": "الدرجة 9",
            "grade_date": "2022-06-01",
            "education_level": "بكالوريوس",
            "specialization": "إعلام واتصال مؤسسي",
            "custom_fields": {
                "phone_number": "091-6789012",
                "residential_address": "طرابلس - النوفليين",
                "blood_type": "O+"
            }
        },
        {
            "national_id": "119830023455",
            "full_name": "خليل مسعود عثمان العبيدي",
            "status_id": status_active,
            "hire_date": "2007-02-14",
            "department": "إدارة التخطيط والمتابعة",
            "current_grade": "الدرجة 12",
            "grade_date": "2020-03-01",
            "education_level": "ماجستير",
            "specialization": "تخطيط استراتيجي",
            "custom_fields": {
                "phone_number": "091-4567890",
                "residential_address": "طرابلس - سوق الجمعة",
                "blood_type": "B+"
            }
        },
        {
            "national_id": "119930034566",
            "full_name": "عبد الرحمن ناصر فرج الكاديكي",
            "status_id": status_active,
            "hire_date": "2017-08-20",
            "department": "إدارة تقنية المعلومات والتوثيق",
            "current_grade": "الدرجة 8",
            "grade_date": "2023-01-10",
            "education_level": "بكالوريوس",
            "specialization": "أمن شبكات ومعلومات",
            "custom_fields": {
                "phone_number": "092-7890123",
                "residential_address": "طرابلس - طريق المطار",
                "blood_type": "A+"
            }
        },
        {
            "national_id": "219860045677",
            "full_name": "سارة كمال بشير الغرياني",
            "status_id": status_active,
            "hire_date": "2010-06-01",
            "department": "قسم الأرشيف والتوثيق",
            "current_grade": "الدرجة 10",
            "grade_date": "2021-08-01",
            "education_level": "ليسانس",
            "specialization": "وثائق ومكتبات",
            "custom_fields": {
                "phone_number": "094-1234567",
                "residential_address": "طرابلس - تاجوراء",
                "blood_type": "O+"
            }
        },
        {
            "national_id": "119980056788",
            "full_name": "حمزة مصطفى خليفة السنوسي",
            "status_id": status_active,
            "hire_date": "2021-02-15",
            "department": "مكتب خدمة المواطنين",
            "current_grade": "الدرجة 6",
            "grade_date": "2021-02-15",
            "education_level": "دبلوم عالي",
            "specialization": "علاقات عامة واستقبال",
            "custom_fields": {
                "phone_number": "091-9012345",
                "residential_address": "طرابلس - حي دمشق",
                "blood_type": "B-"
            }
        },
        {
            "national_id": "119790067899",
            "full_name": "عصام فتحي رمضان البوسيفي",
            "status_id": status_active,
            "hire_date": "2005-04-10",
            "department": "قسم المشتريات والمخازن",
            "current_grade": "الدرجة 11",
            "grade_date": "2020-09-01",
            "education_level": "بكالوريوس",
            "specialization": "إدارة أعمال",
            "custom_fields": {
                "phone_number": "092-2345678",
                "residential_address": "طرابلس - غوط الشعال",
                "blood_type": "A+"
            }
        },
        {
            "national_id": "219940078900",
            "full_name": "هدى محمد إدريس الشريف",
            "status_id": status_active,
            "hire_date": "2018-01-10",
            "department": "إدارة الشؤون الإدارية والمالية",
            "current_grade": "الدرجة 8",
            "grade_date": "2023-05-01",
            "education_level": "بكالوريوس",
            "specialization": "محاسبة وتدقيق",
            "custom_fields": {
                "phone_number": "091-5678901",
                "residential_address": "طرابلس - فشلوم",
                "blood_type": "O+"
            }
        },
        {
            "national_id": "119810089011",
            "full_name": "هيثم مختار ميلاد الشعلالي",
            "status_id": status_active,
            "hire_date": "2006-10-15",
            "department": "المكتب الفني والمشروعات",
            "current_grade": "الدرجة 12",
            "grade_date": "2019-11-01",
            "education_level": "بكالوريوس",
            "specialization": "هندسة كهربائية",
            "custom_fields": {
                "phone_number": "092-8901234",
                "residential_address": "طرابلس - الهضبة الخضراء",
                "blood_type": "AB-"
            }
        },
        {
            "national_id": "219870090122",
            "full_name": "نادية عبد الرؤوف ميلاد الحاسي",
            "status_id": status_leave,
            "hire_date": "2011-08-01",
            "department": "مكتب الموارد البشرية",
            "current_grade": "الدرجة 10",
            "grade_date": "2021-04-01",
            "education_level": "ليسانس",
            "specialization": "علم اجتماع وإدارة",
            "custom_fields": {
                "phone_number": "094-3456789",
                "residential_address": "طرابلس - زناتة",
                "blood_type": "A+"
            }
        },
        {
            "national_id": "119910011233",
            "full_name": "زياد كمال عبد الهادي الرابطي",
            "status_id": status_leave,
            "hire_date": "2015-12-01",
            "department": "قسم العلاقات العامة والإعلام",
            "current_grade": "الدرجة 8",
            "grade_date": "2022-10-01",
            "education_level": "بكالوريوس",
            "specialization": "تصميم جرافيك ووسائط",
            "custom_fields": {
                "phone_number": "091-8765432",
                "residential_address": "طرابلس - جنزور",
                "blood_type": "B+"
            }
        },
        {
            "national_id": "119800022344",
            "full_name": "الصديق عبد الحميد عبد النبي الفرجاني",
            "status_id": status_delegated,
            "hire_date": "2004-03-25",
            "department": "مكتب الشؤون القانونية",
            "current_grade": "الدرجة 12",
            "grade_date": "2018-09-01",
            "education_level": "ماجستير",
            "specialization": "قانون جنائي وإداري",
            "custom_fields": {
                "phone_number": "092-9876543",
                "residential_address": "طرابلس - الحي الإسلامي",
                "blood_type": "O+"
            }
        },
        {
            "national_id": "219930033455",
            "full_name": "نجوى إبراهيم مفتاح الصالحين",
            "status_id": status_delegated,
            "hire_date": "2016-06-15",
            "department": "إدارة التخطيط والمتابعة",
            "current_grade": "الدرجة 8",
            "grade_date": "2023-02-01",
            "education_level": "بكالوريوس",
            "specialization": "إحصاء وبحوث عمليات",
            "custom_fields": {
                "phone_number": "094-5432109",
                "residential_address": "طرابلس - بن غشير",
                "blood_type": "A-"
            }
        },
        {
            "national_id": "119960044566",
            "full_name": "خالد عيسى عاشور الجراري",
            "status_id": status_absent,
            "hire_date": "2019-11-01",
            "department": "مكتب خدمة المواطنين",
            "current_grade": "الدرجة 6",
            "grade_date": "2019-11-01",
            "education_level": "ثانوية عامة",
            "specialization": "قسم أدبي",
            "custom_fields": {
                "phone_number": "091-6543210",
                "residential_address": "طرابلس - صلاح الدين",
                "blood_type": "B+"
            }
        },
        {
            "national_id": "119610055677",
            "full_name": "بشير الصادق عمار التائب",
            "status_id": status_retired,
            "hire_date": "1988-01-15",
            "department": "إدارة الشؤون الإدارية والمالية",
            "current_grade": "الدرجة 14",
            "grade_date": "2014-06-01",
            "education_level": "ليسانس",
            "specialization": "علوم إدارية وقانونية",
            "custom_fields": {
                "phone_number": "092-3210987",
                "residential_address": "طرابلس - قرقارش",
                "blood_type": "O+"
            }
        }
    ]

    inserted_count = 0
    updated_count = 0

    for emp in dummy_employees:
        exists, existing = db.check_national_id_exists(emp["national_id"])
        custom_f = emp.pop("custom_fields", {})
        if exists and existing:
            emp["id"] = existing["id"]
            db.save_employee(emp, custom_f)
            updated_count += 1
        else:
            db.save_employee(emp, custom_f)
            inserted_count += 1

    print(f"[+] اكتمل تلقيم الموظفين: تم إدخال ({inserted_count}) موظف جديد، وتحديث ({updated_count}) موظف.")

    # 3. التأكد من وجود مستخدمين إضافيين لتجربة الصلاحيات وتعدد المستخدمين
    users = {u["display_name"]: u for u in db.get_users_detailed()}
    sample_users = [
        ("مدير النظام", "1234", "admin"),
        ("موظف الإدخال", "0000", "staff"),
        ("أحمد عبد السلام (مدقق)", "1122", "staff"),
        ("مريم الشريف (مسؤول شؤون)", "3344", "staff")
    ]
    for uname, upin, urole in sample_users:
        if uname not in users:
            db.add_user(uname, upin, urole)
            print(f"[+] تم إنشاء المستخدم التجريبي: {uname} (PIN: {upin} - {urole})")

    # 4. توليد ملف إكسل تجريبي للاستيراد وفض التعارض
    generate_sample_import_excel(db)

def generate_sample_import_excel(db: Database) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "سجل الموظفين"
        ws.views.sheetView[0].rightToLeft = True

        headers = [
            "الرقم الوطني",
            "الاسم الكامل",
            "الحالة الوظيفية",
            "تاريخ التعيين",
            "القسم / الإدارة",
            "الدرجة الحالية",
            "تاريخ استحقاق الدرجة",
            "المؤهل العلمي",
            "التخصص الدقيق",
            "رقم الهاتف المحمول",
            "عنوان السكن الحالي"
        ]

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # صفوف تجريبية:
        # السجل 1: موظف جديد تماماً
        # السجل 2: موظف جديد تماماً
        # السجل 3: موظف متعارض (الرقم الوطني 119840012345 مسجل باسم طارق الزوي، ونريد اختبار تحديث قسمه)
        sample_rows = [
            [
                "119970099881",
                "معاذ سالم فتح الله القاضي",
                "على رأس العمل",
                "2022-01-10",
                "مكتب الموارد البشرية",
                "الدرجة 7",
                "2022-01-10",
                "بكالوريوس",
                "إدارة عامة",
                "091-7788990",
                "طرابلس - السراج"
            ],
            [
                "219960088772",
                "ريما خالد عبد الجليل المنفي",
                "على رأس العمل",
                "2021-09-01",
                "إدارة تقنية المعلومات والتوثيق",
                "الدرجة 7",
                "2021-09-01",
                "بكالوريوس",
                "ذكاء اصطناعي",
                "092-4455667",
                "طرابلس - قرجي"
            ],
            [
                "119840012345",
                "طارق عبد السلام محمد الزوي (محدث عبر الاستيراد)",
                "على رأس العمل",
                "2008-04-15",
                "إدارة التحول الرقمي والتقنية",
                "الدرجة 12",
                "2024-01-01",
                "ماجستير",
                "هندسة برمجيات متقدمة",
                "091-3456789",
                "طرابلس - حي الأندلس"
            ]
        ]

        data_font = Font(name="Calibri", size=10)
        data_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D1D5DB"),
            right=Side(style="thin", color="D1D5DB"),
            top=Side(style="thin", color="D1D5DB"),
            bottom=Side(style="thin", color="D1D5DB")
        )

        for row_data in sample_rows:
            ws.append(row_data)
            current_row = ws.max_row
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ملف_استيراد_تجريبي.xlsx")
        wb.save(excel_path)
        print(f"[+] تم إنشاء ملف إكسل تجريبي للاستيراد وفض التعارض بنجاح: {excel_path}")
    except Exception as e:
        print(f"[-] تعذر إنشاء ملف الإكسل التجريبي: {e}")

if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    db_inst = Database()
    seed_database(db_inst)
    print("=== اكتملت عملية التلقيم بنجاح وجاهزية النظام للاختبار الفوري ===")
